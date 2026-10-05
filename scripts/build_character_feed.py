from pathlib import Path
import json
import re

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "data" / "characters"
OUT = ROOT / "data" / "characters.json"

FILES = {
    1: "01_Wuthering_Waves.xlsx",
    2: "02_Genshin_Impact.xlsx",
    3: "03_Honkai_Star_Rail.xlsx",
    4: "04_Arknights_Endfield.xlsx",
    5: "05_Zenless_Zone_Zero.xlsx",
}
GAME_NAMES = {
    1: "Wuthering Waves", 2: "Genshin Impact", 3: "Honkai: Star Rail",
    4: "Arknights: Endfield", 5: "Zenless Zone Zero",
}
GAME_IDS = {1: "wuwa", 2: "genshin", 3: "starrail", 4: "endfield", 5: "zzz"}


def norm(value):
    return re.sub(r"\s+", " ", str(value or "").strip())


def slugify(value):
    s = norm(value).lower().replace("’", "").replace("'", "")
    s = s.replace("•", "-").replace("&", "and")
    return re.sub(r"[^a-z0-9]+", "-", s).strip("-")


def read_table(game_id, path):
    wb = load_workbook(path)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return [], False

    header = [norm(x).lower() for x in rows[0]]
    required = ["id", "имя персонажа", "стихия", "редкость"]
    if header[:4] != required:
        raise SystemExit(f"{path}: expected columns exactly {required}, got {header[:4]}")

    records, changed, used = [], False, set()
    next_number = 1

    for row_index, row in enumerate(rows[1:], start=2):
        if not any(x not in (None, "") for x in row[:4]):
            continue
        name, element, rarity = norm(row[1]), norm(row[2]), row[3]
        if not name:
            raise SystemExit(f"{path}: row {row_index}: character name is empty")

        raw_id = norm(row[0])
        if raw_id:
            if not re.fullmatch(rf"{game_id}\.\d+", raw_id):
                raise SystemExit(f"{path}: row {row_index}: invalid ID {raw_id!r}; expected {game_id}.N")
            char_id = raw_id
            next_number = max(next_number, int(raw_id.split(".", 1)[1]) + 1)
        else:
            while f"{game_id}.{next_number}" in used:
                next_number += 1
            char_id = f"{game_id}.{next_number}"
            ws.cell(row=row_index, column=1, value=char_id)
            changed, next_number = True, next_number + 1

        if char_id in used:
            raise SystemExit(f"{path}: duplicate ID {char_id}")
        used.add(char_id)

        try:
            rarity_int = int(rarity) if rarity not in (None, "") else 0
        except Exception:
            raise SystemExit(f"{path}: row {row_index}: rarity must be numeric")

        records.append({
            "id": char_id,
            "game": GAME_NAMES[game_id],
            "gameId": GAME_IDS[game_id],
            "name": name,
            "slug": slugify(name),
            "element": element,
            "rarity": rarity_int,
            "portrait": f"{char_id}.webp",
        })

    if changed:
        wb.save(path)
    return records, changed


def main():
    SOURCE.mkdir(parents=True, exist_ok=True)
    all_records, changed_files, missing = [], [], []

    for game_id, filename in FILES.items():
        path = SOURCE / filename
        if not path.exists():
            missing.append(filename)
            continue
        records, changed = read_table(game_id, path)
        all_records.extend(records)
        if changed:
            changed_files.append(filename)

    if not all_records:
        print("No Excel character tables found; nothing to generate.")
        return

    ids = [x["id"] for x in all_records]
    if len(ids) != len(set(ids)):
        raise SystemExit("Duplicate character IDs across tables")

    payload = {
        "version": 1,
        "games": [
            {"id": GAME_IDS[gid], "name": GAME_NAMES[gid],
             "characters": [x for x in all_records if x["gameId"] == GAME_IDS[gid]]}
            for gid in range(1, 6)
        ],
    }
    OUT.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"Generated {OUT} with {len(all_records)} characters")
    if missing:
        print("Missing tables:", ", ".join(missing))
    if changed_files:
        print("Assigned IDs in:", ", ".join(changed_files))


if __name__ == "__main__":
    main()

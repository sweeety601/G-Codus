from pathlib import Path
import csv
from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
LIBRARY = ROOT / 'library'
SEED = LIBRARY / 'seed'
FILES = [
    '01_Wuthering_Waves.xlsx',
    '02_Genshin_Impact.xlsx',
    '03_Honkai_Star_Rail.xlsx',
    '04_Arknights_Endfield.xlsx',
    '05_Zenless_Zone_Zero.xlsx',
]


def valid_xlsx(path):
    try:
        with open(path, 'rb') as f:
            if f.read(2) != b'PK':
                return False
        load_workbook(path, read_only=True, data_only=True).close()
        return True
    except Exception:
        return False


def rebuild(path, seed):
    wb = Workbook()
    ws = wb.active
    with open(seed, encoding='utf-8-sig', newline='') as f:
        for row in csv.reader(f):
            ws.append(row)
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    for col, width in {'A': 12, 'B': 34, 'C': 18, 'D': 12}.items():
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=2, max_col=1):
        row[0].number_format = '@'
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


for filename in FILES:
    path = LIBRARY / filename
    seed = SEED / filename.replace('.xlsx', '.csv')
    if not valid_xlsx(path):
        rebuild(path, seed)
        print(f'Rebuilt invalid workbook: {path}')
    else:
        print(f'Workbook is valid: {path}')

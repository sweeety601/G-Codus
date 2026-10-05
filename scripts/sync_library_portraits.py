from pathlib import Path
import json
import re
import shutil
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
LIBRARY = ROOT / "library" / "seed"
IMAGES = ROOT / "images"

TABLES = {
    "01_Wuthering_Waves.xlsx": ("1", "wuwa", "Wuthering Waves"),
    "02_Genshin_Impact.xlsx": ("2", "genshin", "Genshin Impact"),
    "03_Honkai_Star_Rail.xlsx": ("3", "starrail", "Honkai: Star Rail"),
    "04_Arknights_Endfield.xlsx": ("4", "endfield", "Arknights: Endfield"),
    "05_Zenless_Zone_Zero.xlsx": ("5", "zzz", "Zenless Zone Zero"),
}

ALIASES = {
    ("wuwa", "Augusta"): "aug",
    ("wuwa", "Hiyuki"): "hiyuki",
    ("wuwa", "Jianxin"): "jianxin",
    ("wuwa", "Jinhsi"): "jinhsi",
    ("wuwa", "Cartethyia"): "cartethyia",
    ("wuwa", "The Shorekeeper"): "shorekeeper",
    ("wuwa", "Shorekeeper"): "shorekeeper",
    ("wuwa", "Yangyang"): "yangyang",
    ("zzz", "Billy"): "billy-kid",
    ("zzz", "Billy Kid"): "billy-kid",
    ("zzz", "Nicole"): "nicole-demara",
    ("zzz", "Nicole Demara"): "nicole-demara",
    ("zzz", "Anby"): "anby-demara",
    ("zzz", "Anby Demara"): "anby-demara",
    ("zzz", "Grace"): "grace-howard",
    ("zzz", "Grace Howard"): "grace-howard",
    ("zzz", "Koleda"): "koleda",
    ("zzz", "Koleda Belobog"): "koleda",
    ("zzz", "Lucy"): "lucy",
    ("zzz", "Rina"): "rina",
    ("zzz", "Alexandrina Sebastiane"): "rina",
    ("zzz", "Lycaon"): "lycaon",
    ("zzz", "Von Lycaon"): "lycaon",
    ("zzz", "Ellen"): "ellen",
    ("zzz", "Ellen Joe"): "ellen",
    ("zzz", "Jane Doe"): "jane-doe",
    ("zzz", "Miyabi"): "miyabi",
    ("zzz", "Tsukishiro Miyabi"): "miyabi",
    ("zzz", "Harumasa"): "harumasa",
    ("zzz", "Asaba Harumasa"): "harumasa",
    ("zzz", "Yuzuha"): "ukinami-yuzuha",
    ("zzz", "Ukinami Yuzuha"): "ukinami-yuzuha",
    ("zzz", "Seth"): "seth",
    ("zzz", "Seth Lowell"): "seth",
    ("zzz", "Piper"): "piper",
    ("zzz", "Piper Wheel"): "piper",
    ("zzz", "Caesar"): "caesar",
    ("zzz", "Caesar King"): "caesar",
    ("zzz", "Soldier 11"): "soldier-11",
    ("zzz", "Qingyi"): "qingyi",
    ("zzz", "Nekomata"): "nekomata",
    ("zzz", "Nekomiya Mana"): "nekomata",
    ("starrail", "Dan Heng • Imbibitor Lunae"): "imbibitor-lunae",
    ("starrail", "Waveflair"): "aventurine-waveflair",
    ("starrail", "Black"): "black-swan",
    ("starrail", "Swan"): "black-swan",
    ("starrail", "Dan Heng Imbibitor Lunae"): "imbibitor-lunae",
    ("starrail", "Imbibitor Lunae"): "imbibitor-lunae",
    ("starrail", "Mortenax Blade"): "blade-mortenax",
    ("starrail", "March 7th • Evernight"): "march-7th-evernight",
    ("starrail", "March 7th • The Hunt"): "march-7th-swordmaster",
    ("starrail", "Tingyun • Fugue"): "tingyun-fugue",
    ("starrail", "Topaz & Numby"): "topaz",
    ("starrail", "The Herta"): "the-herta",
    ("starrail", "Himeko Nova"): "himeko-nova",
    ("starrail", "Silver Wolf • Lv. 999"): "silver-wolf-lv-999",
    ("endfield", "Last Rite"): "last-rite",
    ("endfield", "Mi Fu"): "mi-fu",
    ("endfield", "Zhuang Fangyi"): "zhuang-fangyi",
}


def norm(s: str) -> str:
    s = s.lower().replace("’", "'")
    return re.sub(r"[^a-z0-9]+", "", s)


def source_files(game_id):
    roots = {
        "genshin": [ROOT / "images_big" / "genshin"],
        "wuwa": [ROOT / "images_big" / "wuthering_waves"],
        "zzz": [ROOT / "images_big" / "zenless_zone_zero"],
        "starrail": [ROOT / "images_big" / "honkai_star_rail", ROOT / "app" / "src" / "main" / "assets" / "honkai_star_rail"],
        "endfield": [ROOT / "images_big" / "arknights_endfield", ROOT / "app" / "src" / "main" / "assets" / "arknights_endfield"],
    }
    files = []
    for root in roots[game_id]:
        if root.exists():
            files.extend(p for p in root.rglob("*") if p.is_file() and p.suffix.lower() in {".webp", ".png", ".jpg", ".jpeg"})
    return files


def stem_for(p):
    stem = p.stem
    for suffix in ("_card", "-card"):
        if stem.endswith(suffix):
            stem = stem[:-len(suffix)]
    return stem


def find_source(game_id, name, files):
    alias = ALIASES.get((game_id, name))
    if alias:
        a = norm(alias)
        exact = [p for p in files if norm(stem_for(p)) == a]
        if exact:
            return exact[0]
    n = norm(name)
    exact = [p for p in files if norm(stem_for(p)) == n]
    if exact:
        return exact[0]
    candidates = [p for p in files if n and (n in norm(stem_for(p)) or norm(stem_for(p)) in n)]
    if len(candidates) == 1:
        return candidates[0]
    return None


def read_rows(path):
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    headers = [str(c.value or "").strip().lower() for c in ws[1]]
    id_col = headers.index("id")
    name_col = headers.index("имя персонажа") if "имя персонажа" in headers else headers.index("имя")
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        name = str(row[name_col] or "").strip()
        cid = str(row[id_col] or "").strip()
        if name and cid:
            rows.append((cid, name))
    return rows


def main():
    IMAGES.mkdir(parents=True, exist_ok=True)
    report = {"copied": [], "missing": [], "duplicates": []}
    seen_ids = set()
    for filename, (prefix, game_id, game_name) in TABLES.items():
        path = LIBRARY / filename
        if not path.exists():
            raise SystemExit(f"Missing library table: {path}")
        files = source_files(game_id)
        for cid, name in read_rows(path):
            if cid in seen_ids:
                report["duplicates"].append(cid)
                continue
            seen_ids.add(cid)
            source = find_source(game_id, name, files)
            if source is None:
                report["missing"].append({"id": cid, "game": game_name, "name": name})
                continue
            target = IMAGES / f"{cid}.webp"
            shutil.copyfile(source, target)
            report["copied"].append({"id": cid, "game": game_name, "name": name, "source": str(source.relative_to(ROOT))})
    (LIBRARY / "generated").mkdir(parents=True, exist_ok=True)
    (LIBRARY / "generated" / "portrait_sync_report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    if report["duplicates"]:
        raise SystemExit(f"Duplicate IDs: {report['duplicates']}")
    print(f"Copied portraits: {len(report['copied'])}")
    print(f"Missing portraits: {len(report['missing'])}")
    for item in report["missing"]:
        print(f"MISSING {item['id']} | {item['game']} | {item['name']}")

if __name__ == "__main__":
    main()

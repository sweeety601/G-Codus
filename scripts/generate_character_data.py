from pathlib import Path
import json
import re
from datetime import datetime, timezone
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
TABLES = {
    "01_Wuthering_Waves.xlsx": ("1", "wuwa", "Wuthering Waves"),
    "02_Genshin_Impact.xlsx": ("2", "genshin", "Genshin Impact"),
    "03_Honkai_Star_Rail.xlsx": ("3", "starrail", "Honkai: Star Rail"),
    "04_Arknights_Endfield.xlsx": ("4", "endfield", "Arknights: Endfield"),
    "05_Zenless_Zone_Zero.xlsx": ("5", "zzz", "Zenless Zone Zero"),
}
OUT = ROOT / "library" / "generated" / "characters.json"
IMAGE_BASE = "https://raw.githubusercontent.com/sweeety601/G-Codus/main/images"
ALIASES = {
    "Имя персонажа": "name", "Имя": "name", "Стихия": "element", "Элемент": "element",
    "Редкость": "rarity", "Rarity": "rarity", "ID": "id",
}


def parse_rarity(v):
    m = re.search(r"\d+", str(v or ""))
    return int(m.group()) if m else 0


def valid_id(value, prefix):
    return bool(re.fullmatch(rf"{re.escape(prefix)}\.\d+", str(value or "").strip()))


def process(path, prefix, game_id, game_name):
    wb = load_workbook(path)
    ws = wb.active
    headers = [str(c.value or "").strip() for c in ws[1]]

    # Portrait is no longer data in Excel. Remove it automatically so old
    # uploaded workbooks become compatible without manual editing.
    portrait_cols = [i + 1 for i, h in enumerate(headers) if h.lower() in {"портрет", "portrait"}]
    for col in reversed(portrait_cols):
        ws.delete_cols(col, 1)
    headers = [str(c.value or "").strip() for c in ws[1]]

    mapped = [ALIASES.get(h, h) for h in headers]
    required = {"name", "element", "rarity"}
    if not required.issubset(set(mapped)):
        raise SystemExit(f"{path}: required columns are Имя персонажа, Стихия, Редкость")

    if "id" not in mapped:
        ws.insert_cols(1)
        ws.cell(1, 1).value = "ID"
        mapped.insert(0, "id")

    id_col = mapped.index("id") + 1
    name_col = mapped.index("name") + 1
    element_col = mapped.index("element") + 1
    rarity_col = mapped.index("rarity") + 1

    rows = []
    for r in range(2, ws.max_row + 1):
        name = str(ws.cell(r, name_col).value or "").strip()
        if not name:
            continue
        rows.append({
            "row": r,
            "id": str(ws.cell(r, id_col).value or "").strip(),
            "name": name,
            "element": str(ws.cell(r, element_col).value or "").strip(),
            "rarity": parse_rarity(ws.cell(r, rarity_col).value),
        })

    used = set()
    for x in rows:
        if x["id"]:
            if not valid_id(x["id"], prefix):
                raise SystemExit(f"{path}: invalid ID {x['id']} at row {x['row']}; expected {prefix}.N")
            number = int(x["id"].split(".", 1)[1])
            if number in used:
                raise SystemExit(f"{path}: duplicate ID {x['id']} at row {x['row']}")
            used.add(number)

    counter = max(used, default=0) + 1
    for x in rows:
        if not x["id"]:
            x["id"] = f"{prefix}.{counter}"
            ws.cell(x["row"], id_col).number_format = "@"
            ws.cell(x["row"], id_col).value = x["id"]
            used.add(counter)
            counter += 1

    result = [{
        "id": x["id"], "gameId": game_id, "name": x["name"],
        "element": x["element"], "rarity": x["rarity"], "announced": False,
        "portraitUrl": f"{IMAGE_BASE}/{x['id']}.webp"
    } for x in rows]
    wb.save(path)
    return {"id": prefix, "gameId": game_id, "name": game_name, "characters": result}


def main():
    base = ROOT / "library" / "seed"
    games = []
    for filename, meta in TABLES.items():
        path = base / filename
        if not path.exists():
            raise SystemExit(f"Missing source: {path}")
        games.append(process(path, *meta))
    payload = {"version": 1, "generatedAt": datetime.now(timezone.utc).isoformat(), "games": games}
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


if __name__ == "__main__":
    main()

from pathlib import Path
from openpyxl import load_workbook

ROOT=Path(".")
SEED=ROOT/"library"/"seed"
BANNERS=ROOT/"banners"

CONFIG=[
 ("01_Wuthering_Waves","3.7 Phase 2","2026-10-22","2026-11-11",["Suoming","Lynae","Lucilla"],[]),
 ("02_Genshin_Impact","7.1 Phase 2","2026-10-13","2026-11-03",["Skirk","Escoffier"],[]),
 ("03_Honkai_Star_Rail","4.6 Phase 2","2026-10-21","2026-11-10",["Pearl","Mortenax Blade"],["Misha","Xueyi","Qingque"]),
 ("04_Arknights_Endfield","1.6 Phase 1","2026-10-15","2026-11-05",["Si"],[]),
 ("05_Zenless_Zone_Zero","3.3 Phase 1","2026-10-21","2026-11-11",["Phoenix"],[]),
]

def n(v):
    return " ".join(str(v or "").strip().lower().replace("ё","е").split())

def seed_map(base, wanted):
    wb=load_workbook(SEED/(base+".xlsx"),read_only=True,data_only=True)
    ws=wb.active
    rows=list(ws.iter_rows(values_only=True))
    wb.close()
    out={}
    for wanted_name in wanted:
        key=n(wanted_name)
        matches=[]
        for row in rows[1:]:
            cells=[str(v or "").strip() for v in row]
            if not any(n(v)==key for v in cells):
                continue
            ids=[v for v in cells if __import__("re").match(r"^\\d+\\.\\d+$",v)]
            if len(ids)==1:
                matches.append(ids[0])
        if len(set(matches))!=1:
            raise RuntimeError(f"Cannot uniquely map {wanted_name} in {base}: {matches[:10]}")
        out[wanted_name]=matches[0]
    return out

def col(headers, aliases, required=True):
    aliases={n(x) for x in aliases}
    for i,h in enumerate(headers):
        if h in aliases:return i+1
    if required: raise RuntimeError(f"Missing column {aliases}; headers={headers}")
    return None

def update(base,phase,start,end,five,four):
    path=BANNERS/f"{base}_confirmed.xlsx"
    wb=load_workbook(path)
    ws=wb.active
    headers=[n(c.value) for c in ws[1]]
    pc=col(headers,["версия и фаза","phase","version and phase","версия","version"])
    sc=col(headers,["дата начала","start date","start_date","начало","start"])
    ec=col(headers,["дата окончания","end date","end_date","конец","end"])
    cc=col(headers,["персонажи в составе баннера","персонажи","characters","character","featured characters"],False)
    fc=col(headers,["5* в баннере","5★ в баннере","5*","5 star","5-star","five star","five_star","5star","featured 5 star","featured 5★","персонажи 5*","персонажи 5★"],False)
    qc=col(headers,["4* в баннере","4★ в баннере","4*","4 star","4-star","four star","four_star","4star","featured 4 star","featured 4★","персонажи 4*","персонажи 4★"],False)
    ids=seed_map(base,five+four)
    five_ids=[ids[x] for x in five]
    four_ids=[ids[x] for x in four]
    row=next((r for r in range(2,ws.max_row+1) if n(ws.cell(r,pc).value)==n(phase)),ws.max_row+1)
    ws.cell(row,pc).value=phase; ws.cell(row,sc).value=start; ws.cell(row,ec).value=end
    if fc: ws.cell(row,fc).value="; ".join(five_ids)
    if qc: ws.cell(row,qc).value="; ".join(four_ids)
    if cc: ws.cell(row,cc).value="; ".join(five_ids)
    wb.save(path); wb.close()
    print(base,phase,{**ids})

for cfg in CONFIG: update(*cfg)

# Lynae/Lucilla are now confirmed: remove their IDs/names from WuWa 3.7 Phase 2 leak row.
base="01_Wuthering_Waves"
leak=BANNERS/f"{base}_leaks.xlsx"
m=seed_map(base,["Lynae","Lucilla"])
wb=load_workbook(leak); ws=wb.active; headers=[n(c.value) for c in ws[1]]
pc=col(headers,["версия и фаза","phase","version and phase","версия","version"])
cols=[col(headers,["персонажи в составе баннера","персонажи","characters","character","featured characters"],False),
      col(headers,["5* в баннере","5★ в баннере","5*","5 star","5-star","five star","five_star","5star","featured 5 star","featured 5★","персонажи 5*","персонажи 5★"],False),
      col(headers,["4* в баннере","4★ в баннере","4*","4 star","4-star","four star","four_star","4star","featured 4 star","featured 4★","персонажи 4*","персонажи 4★"],False)]
remove={n("Lynae"),n("Lucilla"),m["Lynae"],m["Lucilla"]}
for r in range(2,ws.max_row+1):
    if n(ws.cell(r,pc).value)!="3.7 phase 2": continue
    for c in cols:
        if not c: continue
        vals=[x.strip() for x in str(ws.cell(r,c).value or "").replace("|",";").split(";") if x.strip()]
        ws.cell(r,c).value="; ".join(x for x in vals if n(x) not in remove)
wb.save(leak); wb.close()
print("Removed confirmed Lynae/Lucilla from WuWa leak Phase 2.")

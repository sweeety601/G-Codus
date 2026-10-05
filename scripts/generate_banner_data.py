from pathlib import Path
import json
import re
from datetime import datetime, timezone
from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
BANNERS = ROOT / "banners"
OUT = ROOT / "data" / "banner_feed.json"

GAMES = {
    "01_Wuthering_Waves": ("wuwa", "Wuthering Waves"),
    "02_Genshin_Impact": ("genshin", "Genshin Impact"),
    "03_Honkai_Star_Rail": ("starrail", "Honkai: Star Rail"),
    "04_Arknights_Endfield": ("endfield", "Arknights: Endfield"),
    "05_Zenless_Zone_Zero": ("zzz", "Zenless Zone Zero"),
}

HEADERS = ["Версия и фаза", "Дата начала", "Дата окончания", "Персонажи в составе баннера", "4* в баннере"]


def split_names(value):
    if value is None:
        return []
    return [x.strip() for x in re.split(r"[,;\\n]+", str(value)) if x.strip()]


def make_initial_workbook(path):
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Banners"
    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = cell.font.copy(bold=True)
    ws.freeze_panes = "A2"
    widths = [24, 16, 16, 42, 36]
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i)].width = width
    wb.save(path)


def read_table(path):
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    headers = [str(c.value or "").strip().lower() for c in ws[1]]
    wanted = [h.lower() for h in HEADERS]
    indexes = {}
    for h in wanted:
        if h not in headers:
            raise SystemExit(f"{path}: missing column '{h}'")
        indexes[h] = headers.index(h)

    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        phase = str(row[indexes[HEADERS[0].lower()]] or "").strip()
        if not phase:
            continue
        if not re.fullmatch(r"\d+\.\d+ Phase \d+", phase, re.IGNORECASE):
            raise SystemExit(f"{path}: invalid version/phase '{phase}'")
        start = str(row[indexes[HEADERS[1].lower()]] or "").strip()
        end = str(row[indexes[HEADERS[2].lower()]] or "").strip()
        characters = split_names(row[indexes[HEADERS[3].lower()]])
        four = split_names(row[indexes[HEADERS[4].lower()]])
        if not start or not end:
            raise SystemExit(f"{path}: {phase} requires start and end dates")
        rows.append({
            "phase": phase,
            "start": start,
            "end": end,
            "characters": characters,
            "four_star": four,
        })
    return rows


def parse_date(value, end=False):
    value = value.strip()
    for fmt in ("%Y-%m-%d", "%Y-%m-%d %H:%M", "%Y-%m-%dT%H:%M:%S", "%Y-%m-%dT%H:%M:%SZ"):
        try:
            dt = datetime.strptime(value, fmt)
            if dt.tzinfo is None:
                dt = dt.replace(tzinfo=timezone.utc)
            return dt
        except ValueError:
            pass
    raise SystemExit(f"Invalid date: {value}")


def normalize(row, confirmed):
    return {
        "phase": row["phase"],
        "start": row["start"],
        "end": row["end"],
        "characters": row["characters"],
        "four_star": row["four_star"],
        "source_status": "confirmed" if confirmed else "unconfirmed",
        "unconfirmed": not confirmed,
    }


def build_game(game_key, game_name):
    confirmed_path = BANNERS / f"{game_key}_confirmed.xlsx"
    leaks_path = BANNERS / f"{game_key}_leaks.xlsx"
    if not confirmed_path.exists():
        make_initial_workbook(confirmed_path)
    if not leaks_path.exists():
        make_initial_workbook(leaks_path)

    confirmed = {r["phase"].lower(): normalize(r, True) for r in read_table(confirmed_path)}
    leaks = {r["phase"].lower(): normalize(r, False) for r in read_table(leaks_path)}

    merged = dict(leaks)
    merged.update(confirmed)
    rows = list(merged.values())
    rows.sort(key=lambda x: parse_date(x["start"]))

    now = datetime.now(timezone.utc)
    current = [x for x in rows if parse_date(x["start"]) <= now < parse_date(x["end"], True)]
    future = [x for x in rows if parse_date(x["start"]) > now]
    history = [x for x in rows if parse_date(x["end"], True) <= now]

    # The app displays the immediate next phase separately and everything after it
    # as upcoming. Each row retains its own confirmed/unconfirmed status.
    next_rows = future[:1]
    upcoming = future[1:]
    return {
        "id": game_key,
        "name": game_name,
        "current": current,
        "next": next_rows,
        "upcoming": upcoming,
        "history": list(reversed(history)),
        "source": "G-Codus banner seed tables",
    }


def main():
    BANNERS.mkdir(parents=True, exist_ok=True)
    games = [build_game(key, meta[1]) for key, meta in GAMES.items()]
    payload = {
        "version": 4,
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "source": "G-Codus banner Excel tables",
        "games": {g["name"]: g for g in games},
    }
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print("Generated banner feed from 10 Excel tables.")
    for g in games:
        print(g["name"], "current=", len(g["current"]), "next=", len(g["next"]), "upcoming=", len(g["upcoming"]), "history=", len(g["history"]))


if __name__ == "__main__":
    main()
